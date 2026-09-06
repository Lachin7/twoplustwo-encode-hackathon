"""Shared by llm_predict.py and tinker_predict.py: prompt, answer schema, output files, traces."""

import asyncio
import datetime
import json
import os
import re
import shutil
import sys
import threading
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.cell.cell import MergedCell
from pydantic import BaseModel, Field

from sb import answer_cells, load_dataset, read_jsonl, serialize_workbook

ENV_PATH = Path(__file__).resolve().parent.parent / ".env"

SYSTEM_PROMPT = (
    "You are a spreadsheet expert. You get a serialized workbook and a user instruction. "
    "Fill the answer range. Literals and Excel formulas (strings starting with =) are both fine. "
    "For a long column or row, one relative formula on the first cell is enough — the harness "
    "fills the rest and LibreOffice recalculates. Use null for cells that must be empty."
)
FORMAT_HINT = (
    '\n\nReply with JSON only, no prose, in this shape: '
    '{"cells": [{"cell": "B6", "value": "=A6*2"}, {"cell": "C2", "value": 42}]}'
)
RETRY_HINT = (
    "\n\nYour previous reply was not valid JSON. "
    "Reply with JSON only, no prose: "
    '{"cells": [{"cell": "B6", "value": "=A6*2"}]}'
)
THIN_HINT = (
    "\n\nThat JSON covered too few answer cells. "
    "Return one relative formula on the first cell of each column or row, or every cell."
)

_IO_LOCK = threading.Lock()


class CellValue(BaseModel):
    cell: str = Field(description="Cell address like A3 or B6")
    value: str | int | float | bool | None = Field(description="Final value for that cell")


class SpreadsheetAnswer(BaseModel):
    cells: list[CellValue]


def load_env(path: Path = ENV_PATH) -> None:
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ[key.strip()] = value.strip().strip('"').strip("'")


def parse_ids(text: str | None) -> set[str] | None:
    return {i.strip() for i in text.split(",") if i.strip()} if text else None


def selected_tasks(dataset_dir: Path, ids: set[str] | None) -> list[dict]:
    tasks = load_dataset(dataset_dir)
    return tasks if ids is None else [t for t in tasks if t["id"] in ids]


def graded_count(task: dict, wb=None) -> int:
    own = wb is None
    if own:
        wb = openpyxl.load_workbook(task["init_xlsx"], data_only=True)
    try:
        return len(answer_cells(task, wb))
    finally:
        if own:
            wb.close()


def build_prompt(task: dict) -> str:
    n = graded_count(task)
    # Agent tasks can inspect the live workbook. A smaller preview prevents wide,
    # multi-sheet workbooks from exceeding Qwen's context window before turn 1.
    workbook = serialize_workbook(
        task["init_xlsx"],
        max_rows=40 if n > 20 else 120,
        max_cols=30,
        task=task,
    )
    if len(workbook) > 100_000:
        workbook = workbook[:100_000] + "\n\n[Preview truncated; inspect wb directly with Python.]"
    strategy = (
        f"{n} graded cells — one relative formula per column or row on the first answer cell, not a value dump."
        if n > 20
        else f"{n} graded cells — literals or short formulas are both fine."
    )
    return (
        f"## Instruction\n{task['instruction']}\n\n"
        f"## Workbook\n{workbook}\n\n"
        f"## Answer range\nSheet: {task.get('answer_sheet') or 'active sheet'}\n"
        f"Cells: {task['answer_position']}\n{strategy}\n"
    )


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
    "%d-%b-%Y",
    "%d-%b-%y",
)
_TIME_FORMATS = ("%H:%M:%S", "%H:%M")


def parse_writable_value(value, number_format: str = "General"):
    """Turn date/time strings (and duration formats) into Excel-native types."""
    fmt = (number_format or "General").lower()
    duration = "[h]" in fmt or "[hh]" in fmt
    if isinstance(value, datetime.datetime) and value.year in (1899, 1900) and value.date() <= datetime.date(1900, 1, 1):
        if value.time() != datetime.time():
            return datetime.timedelta(hours=value.hour, minutes=value.minute, seconds=value.second) if duration else value.time()
    if not isinstance(value, str):
        if duration and isinstance(value, (int, float)) and 0 <= float(value) < 10:
            return datetime.timedelta(days=float(value))
        if ("h" in fmt or "mm" in fmt) and "y" not in fmt and isinstance(value, (int, float)) and 0 < float(value) < 1:
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(value))).time()
        return value
    if value.lstrip().startswith("="):
        return normalize_formula(value)
    text = value.strip()
    if not text:
        return None
    if text.lower().startswith("1 day, "):
        text = text.split(",", 1)[1].strip()
        try:
            clock = datetime.datetime.strptime(text, "%H:%M:%S")
            return datetime.timedelta(days=1, hours=clock.hour, minutes=clock.minute, seconds=clock.second)
        except ValueError:
            pass
    for fmt_s in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt_s)
        except ValueError:
            continue
    for fmt_s in _TIME_FORMATS:
        try:
            clock = datetime.datetime.strptime(text, fmt_s)
            delta = datetime.timedelta(hours=clock.hour, minutes=clock.minute, seconds=clock.second)
            return delta if duration else clock.time()
        except ValueError:
            continue
    return value


def coerce_cell_value(value: str | int | float | bool | None) -> str | int | float | bool | None:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "" or stripped.lower() in ("null", "none"):
            return None
        if stripped.startswith("="):
            return stripped
        if stripped.isdigit() and stripped.startswith("0") and stripped != "0":
            return stripped
        try:
            number = float(stripped.replace(",", ""))
            return int(number) if number.is_integer() else number
        except ValueError:
            return value
    return value


_XLFN = (
    "XLOOKUP", "XMATCH", "UNIQUE", "SORT", "SORTBY", "SEQUENCE",
    "LET", "LAMBDA", "CHOOSECOLS", "CHOOSEROWS", "HSTACK", "VSTACK",
    "TOCOL", "TOROW", "TAKE", "DROP", "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER",
    "WRAPROWS", "WRAPCOLS",
)


def normalize_formula(value):
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text.startswith("="):
        return value
    text = re.sub(r"(?<!_xlfn\._xlws\.)\bFILTER\s*\(", "_xlfn._xlws.FILTER(", text, flags=re.I)
    for name in _XLFN:
        text = re.sub(rf"(?<!_xlfn\.)\b{name}\s*\(", rf"_xlfn.{name}(", text, flags=re.I)
    return text


def _col_index(col: str) -> int:
    n = 0
    for ch in col.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _col_letters(index: int) -> str:
    out = ""
    n = index
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _shift_formula(formula: str, row_delta: int = 0, col_delta: int = 0) -> str:
    if row_delta == 0 and col_delta == 0:
        return formula

    def repl(m: re.Match) -> str:
        abs_col, col, abs_row, row = m.group(1), m.group(2), m.group(3), m.group(4)
        new_col = col if abs_col or col_delta == 0 else _col_letters(_col_index(col) + col_delta)
        new_row = row if abs_row or row_delta == 0 else str(int(row) + row_delta)
        return f"{abs_col}{new_col}{abs_row}{new_row}"

    return re.sub(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)", repl, formula)


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("=")


def parse_answer(text: str) -> SpreadsheetAnswer:
    """JSON object from a reply. Thinking models wrap it in prose, fences, or extra tokens."""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    text = re.sub(r"```(?:json)?\s*", "", text, flags=re.I)
    start = text.find("{")
    if start < 0:
        raise ValueError(f"no JSON object in reply: {text[:120]!r}")
    snippet = text[start:]
    try:
        payload, _ = json.JSONDecoder().raw_decode(snippet)
        return SpreadsheetAnswer.model_validate(payload)
    except json.JSONDecodeError:
        end = snippet.rfind("}")
        if end < 0:
            raise ValueError(f"no JSON object in reply: {text[:120]!r}") from None
        repaired = re.sub(r",\s*([}\]])", r"\1", snippet[: end + 1])
        return SpreadsheetAnswer.model_validate(json.loads(repaired))


def _split_cell(cell: str) -> tuple[str | None, str]:
    text = cell.strip().replace("'", "")
    if "!" in text:
        sheet, coord = text.rsplit("!", 1)
        return sheet, coord.upper()
    return None, text.upper()


def _sheet_key(sheet: str | None) -> str:
    return (sheet or "").strip().upper()


def align_answer(task: dict, answer: SpreadsheetAnswer, wb=None) -> SpreadsheetAnswer:
    """Map model cells onto the graded range. Prefer Sheet!A1; zip only if keys are bare."""
    expected = [(sheet, coord.upper()) for sheet, coord in answer_cells(task, wb)]
    raw = [(*_split_cell(cell.cell), coerce_cell_value(cell.value)) for cell in answer.cells]
    coords = [coord for _sheet, coord in expected]
    unique = len(coords) == len(set(coords))
    by_sheet = {(_sheet_key(sheet), coord): value for sheet, coord, value in raw if sheet}
    by_coord = {coord: value for _sheet, coord, value in raw}
    if expected and all((_sheet_key(sheet), coord) in by_sheet for sheet, coord in expected):
        cells = [CellValue(cell=coord, value=by_sheet[(_sheet_key(sheet), coord)]) for sheet, coord in expected]
    elif unique and expected and all(coord in by_coord for coord in coords):
        cells = [CellValue(cell=coord, value=by_coord[coord]) for coord in coords]
    elif expected and raw:
        cells = [
            CellValue(cell=coord, value=raw[i][2])
            for i, (_sheet, coord) in enumerate(expected)
            if i < len(raw)
        ]
    else:
        cells = [CellValue(cell=coord, value=value) for _sheet, coord, value in raw]
    return SpreadsheetAnswer(cells=cells)


def set_cell_value(ws, coord: str, value) -> None:
    cell = ws[coord]
    target = cell
    if isinstance(cell, MergedCell):
        target = None
        for merged in ws.merged_cells.ranges:
            if coord in merged:
                target = ws.cell(merged.min_row, merged.min_col)
                break
        if target is None:
            return
    fmt = getattr(target, "number_format", "General")
    target.value = parse_writable_value(value, fmt)


def _col_row(coord: str) -> tuple[str, int]:
    col = "".join(ch for ch in coord if ch.isalpha())
    row = int("".join(ch for ch in coord if ch.isdigit()))
    return col, row


def _expand_line(
    out: dict[tuple[str | None, str], object],
    sheet: str | None,
    cells: list[tuple[str, int, int]],
    axis: str,
) -> None:
    """cells are (coord, row, col_index). Exactly one formula seed fills the line."""

    def key(coord: str) -> tuple[str | None, str]:
        if (sheet, coord) in out:
            return (sheet, coord)
        if (None, coord) in out:
            return (None, coord)
        return (sheet, coord)

    seeds = [(coord, row, col) for coord, row, col in cells if _is_formula(out.get(key(coord)))]
    if len(seeds) != 1:
        return
    seed, seed_row, seed_col = seeds[0]
    formula = normalize_formula(out[key(seed)])
    overwrite = len(cells) > 20
    for coord, row, col in cells:
        if coord == seed:
            continue
        k = key(coord)
        cur = out.get(k)
        if k in out and cur is not None and not (overwrite and not _is_formula(cur)):
            continue
        out[k] = _shift_formula(formula, row - seed_row, 0) if axis == "col" else _shift_formula(formula, 0, col - seed_col)


def expand_formulas(expected: list[tuple[str, str]], given: dict) -> dict:
    """Fill the dominant axis from a single relative formula. Overwrite junk only on long lines.

    Tall ranges fill down only. Wide ranges fill right only. Never both — inventing a
    second column from the first on a 2-col block is a common way to fail a pass.
    Keys may be coord strings or (sheet, coord) tuples.
    """
    if not given:
        return given
    tuple_in = isinstance(next(iter(given)), tuple)
    out: dict[tuple[str | None, str], object] = dict(given) if tuple_in else {(None, k): v for k, v in given.items()}
    by_col: dict[tuple[str | None, str], list[tuple[str, int, int]]] = {}
    by_row: dict[tuple[str | None, int], list[tuple[str, int, int]]] = {}
    for sheet, coord in expected:
        col, row = _col_row(coord)
        item = (coord, row, _col_index(col))
        by_col.setdefault((sheet, col), []).append(item)
        by_row.setdefault((sheet, row), []).append(item)
    sheets = {sheet for sheet, _ in by_col} | {sheet for sheet, _ in by_row}
    for sheet in sheets:
        col_groups = [cells for (s, _), cells in by_col.items() if s == sheet]
        row_groups = [cells for (s, _), cells in by_row.items() if s == sheet]
        max_col = max((len(g) for g in col_groups), default=0)
        max_row = max((len(g) for g in row_groups), default=0)
        if max_col >= max_row:
            for cells in col_groups:
                cells.sort(key=lambda item: item[1])
                _expand_line(out, sheet, cells, "col")
        else:
            for cells in row_groups:
                cells.sort(key=lambda item: item[2])
                _expand_line(out, sheet, cells, "row")
    return out if tuple_in else {coord: value for (_sheet, coord), value in out.items()}


def output_is_thin(task: dict, out_path: Path) -> bool:
    """True when a large range was barely written. Small tasks never retry — they already pass."""
    if not out_path.exists():
        return True
    pred = openpyxl.load_workbook(out_path)
    init = openpyxl.load_workbook(task["init_xlsx"])
    try:
        expected = list(answer_cells(task, pred))
        if len(expected) <= 20:
            return False
        wrote = 0
        for sheet, coord in expected:
            pws = pred[sheet] if sheet and sheet in pred.sheetnames else pred.active
            iws = init[sheet] if sheet and sheet in init.sheetnames else init.active
            pv, iv = pws[coord].value, iws[coord].value
            if _is_formula(pv) or (pv is not None and pv != iv):
                wrote += 1
        return wrote < max(2, int(0.35 * len(expected)))
    finally:
        pred.close()
        init.close()


def write_output(task: dict, answer: SpreadsheetAnswer, out_path: Path) -> None:
    shutil.copy(task["init_xlsx"], out_path)
    wb = openpyxl.load_workbook(out_path)
    aligned = align_answer(task, answer, wb)
    expected = [(sheet, coord.upper()) for sheet, coord in answer_cells(task, wb)]
    cells = {}
    for i, (sheet, coord) in enumerate(expected):
        if i >= len(aligned.cells):
            break
        value = aligned.cells[i].value
        cells[(sheet, coord)] = normalize_formula(value) if isinstance(value, str) else value
    cells = expand_formulas(expected, cells)
    for sheet, coord in expected:
        key = (sheet, coord)
        if key not in cells:
            continue
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        set_cell_value(ws, coord, cells[key])
    wb.save(out_path)


def prepare_out_dir(out_dir: Path) -> None:
    for sub in ("outputs", "traces"):
        shutil.rmtree(out_dir / sub, ignore_errors=True)
        (out_dir / sub).mkdir(parents=True)
    for name in ("predictions.jsonl", "run.log"):
        (out_dir / name).write_text("", encoding="utf-8")


def append_jsonl(path: Path, record: dict) -> None:
    line = json.dumps(record, ensure_ascii=False) + "\n"
    with _IO_LOCK:
        with path.open("a", encoding="utf-8") as f:
            f.write(line)


def log(out_dir: Path, line: str) -> None:
    print(line, flush=True)
    with _IO_LOCK:
        with (out_dir / "run.log").open("a", encoding="utf-8") as f:
            f.write(line + "\n")


async def predict_task(
    complete,
    model: str,
    task: dict,
    out_dir: Path,
    *,
    record_prediction: bool = True,
) -> str:
    """One model call, retry on bad JSON or a too-thin fill. On total failure copy the init workbook.

    `complete(prompt)` is an async function returning (text, input_tokens, output_tokens).
    """
    out = out_dir / "outputs" / f"{task['id']}.xlsx"
    prompt = build_prompt(task)
    status = "ok"
    last_error: Exception | None = None
    extra = ""
    wrote_ok = False
    for step in range(1, 4):
        trace = {"step": step, "model": model, "prompt": None, "response": None,
                 "input_tokens": None, "output_tokens": None, "latency_ms": None, "error": None}
        started = time.time()
        try:
            trace["prompt"] = prompt + extra
            text, trace["input_tokens"], trace["output_tokens"] = await complete(trace["prompt"])
            trace["response"] = text
            write_output(task, parse_answer(text), out)
            last_error = None
            wrote_ok = True
            status = "ok"
            if output_is_thin(task, out):
                trace["error"] = "thin fill"
                extra = THIN_HINT
            else:
                extra = ""
        except Exception as e:
            last_error = e
            trace["error"] = f"{type(e).__name__}: {e}"[:500]
            status = f"error: {e}"[:200]
            extra = RETRY_HINT
        trace["latency_ms"] = int((time.time() - started) * 1000)
        append_jsonl(out_dir / "traces" / f"{task['id']}.jsonl", trace)
        if last_error is None and extra == "":
            break
    if last_error is not None and not wrote_ok:
        shutil.copy(task["init_xlsx"], out)
    if record_prediction:
        append_jsonl(
            out_dir / "predictions.jsonl",
            {"id": task["id"], "output": f"outputs/{task['id']}.xlsx", "status": status},
        )
    return status


async def run(
    complete,
    model: str,
    tasks: list[dict],
    out_dir: Path,
    concurrency: int,
    resume: bool = False,
    agent: str = "auto",
    complete_chat=None,
) -> None:
    if resume and (out_dir / "predictions.jsonl").exists():
        done = {row["id"] for row in read_jsonl(out_dir / "predictions.jsonl")}
        tasks = [task for task in tasks if task["id"] not in done]
        (out_dir / "outputs").mkdir(parents=True, exist_ok=True)
        (out_dir / "traces").mkdir(parents=True, exist_ok=True)
        log(out_dir, f"resume skip {len(done)}  remaining {len(tasks)}")
    else:
        prepare_out_dir(out_dir)
    log(out_dir, f"model {model}  tasks {len(tasks)}  agent={agent}")
    if not tasks:
        return
    semaphore = asyncio.Semaphore(concurrency)

    async def run_one(task: dict) -> None:
        try:
            async with semaphore:
                if agent == "off":
                    status = await predict_task(complete, model, task, out_dir)
                else:
                    from agent import predict_routed

                    status = await predict_routed(complete, complete_chat, model, task, out_dir, agent)
        except Exception as e:
            status = f"error: {type(e).__name__}: {e}"[:200]
            append_jsonl(
                out_dir / "predictions.jsonl",
                {"id": task["id"], "output": f"outputs/{task['id']}.xlsx", "status": status},
            )
        log(out_dir, f"{task['id']:<8} {status}")

    await asyncio.gather(*(run_one(task) for task in tasks))
