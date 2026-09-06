"""Isolated checks for the formula / JSON-parse harness changes.

    uv run baseline/test_harness_changes.py
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from common import (
    _shift_formula,
    expand_formulas,
    normalize_formula,
    output_is_thin,
    parse_answer,
    write_output,
)
from evaluate import score_task
from sb import load_dataset, recalculate, serialize_workbook, soffice_path

PASS = 0
FAIL = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")


def old_parse(text: str):
    """Pre-change parser: first {{ to last }}, no trailing-comma / raw_decode."""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise ValueError("no JSON")
    return json.loads(text[start : end + 1])


def test_json_parse() -> None:
    print("\n[1] thinking → JSON (vs old parser)")
    trailing_comma = '{"cells":[{"cell":"B6","value":4},]}'
    extra_brace = '{"cells":[{"cell":"B6","value":4}]} see also {more}'
    fences = '<think>plan {bad}</think>\n```json\n{"cells":[{"cell":"B6","value":4}]}\n```'
    prose = 'Here you go\n{"cells":[{"cell":"B6","value":"=A6*2"}]} extra'

    cases = [
        ("trailing comma", trailing_comma, True),
        ("extra { after object", extra_brace, True),
        ("think + fence", fences, False),
        ("prose wrap", prose, False),
    ]
    for name, raw, old_should_fail in cases:
        new_ok = True
        try:
            parse_answer(raw)
        except Exception as e:
            new_ok = False
            check(f"new parses {name}", False, str(e))
            continue
        old_ok = True
        try:
            old_parse(raw)
        except Exception:
            old_ok = False
        if old_should_fail:
            check(f"old fails {name} (so new is the fix)", not old_ok)
        check(f"new parses {name}", new_ok)


def test_xlfn() -> None:
    print("\n[2] _xlfn. prefix")
    check("XLOOKUP", normalize_formula("=XLOOKUP(A1,B:B,C:C)") == "=_xlfn.XLOOKUP(A1,B:B,C:C)")
    check("UNIQUE", "_xlfn.UNIQUE" in normalize_formula("=UNIQUE(A1:A10)"))
    check("FILTER uses _xlws", "_xlfn._xlws.FILTER" in normalize_formula("=FILTER(A1:B10,A1:A10>0)"))
    check("already prefixed stays", normalize_formula("=_xlfn.XLOOKUP(A1,B:B,C:C)") == "=_xlfn.XLOOKUP(A1,B:B,C:C)")
    check("literals unchanged", normalize_formula("hello") == "hello")


def test_filldown() -> None:
    print("\n[3] formula fill-down")
    check("relative row", _shift_formula("=A2*2+$B$1", 3) == "=A5*2+$B$1")
    expected = [(None, "C2"), (None, "C3"), (None, "C4")]
    out = expand_formulas(expected, {"C2": "=A2+1"})
    check("fills C3", out.get("C3") == "=A3+1", str(out))
    check("fills C4", out.get("C4") == "=A4+1", str(out))
    no = expand_formulas(expected, {"C2": 10, "C3": 11, "C4": 12})
    check("values not rewritten", no == {"C2": 10, "C3": 11, "C4": 12})
    check("abs col stays on right-shift", _shift_formula("=$A2+B2", 0, 1) == "=$A2+C2")
    row = [(None, "C5"), (None, "D5"), (None, "E5")]
    across = expand_formulas(row, {"C5": "=A5*2"})
    check("fills D5", across.get("D5") == "=B5*2", str(across))
    check("fills E5", across.get("E5") == "=C5*2", str(across))
    tall2 = [(None, f"{col}{r}") for r in range(1, 13) for col in ("A", "B")]
    no_invent = expand_formulas(tall2, {"A1": "=C1*2"})
    check("tall 2-col does not invent B from A", "B1" not in no_invent, str(no_invent))
    check("tall 2-col still fills A", no_invent.get("A12") == "=C12*2", str(no_invent))
    long_col = [(None, f"H{r}") for r in range(2, 42)]
    junk = {f"H{r}": 0 for r in range(2, 42)}
    junk["H2"] = "=G2+1"
    filled = expand_formulas(long_col, junk)
    check("large col overwrites junk", filled.get("H40") == "=G40+1", str(filled.get("H40")))
    short_junk = expand_formulas(expected, {"C2": "=A2+1", "C3": 99})
    check("small col keeps literal", short_junk.get("C3") == 99, str(short_junk))


def test_filldown_recalc() -> None:
    print("\n[4] fill-down + LibreOffice (synthetic column)")
    if not soffice_path():
        check("soffice present", False, "skipped")
        return
    tmp = Path(tempfile.mkdtemp())
    init = tmp / "init.xlsx"
    gold = tmp / "golden.xlsx"
    pred = tmp / "pred.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i, n in enumerate((1, 2, 3, 4, 5), start=1):
        ws[f"A{i}"] = n
        ws[f"B{i}"] = n * 2
    wb.save(gold)
    for i in range(1, 6):
        ws[f"B{i}"] = None
    wb.save(init)
    task = {
        "id": "synth-fill",
        "instruction": "double column A into B",
        "init_xlsx": str(init),
        "golden_xlsx": str(gold),
        "answer_position": "B1:B5",
        "answer_sheet": "Sheet1",
        "instruction_type": "Cell-Level Manipulation",
    }
    from common import SpreadsheetAnswer, CellValue

    one = SpreadsheetAnswer(cells=[CellValue(cell="B1", value="=A1*2")])
    write_output(task, one, pred)
    written = openpyxl.load_workbook(pred)["Sheet1"]
    check("B5 got a formula", isinstance(written["B5"].value, str) and written["B5"].value.startswith("="), str(written["B5"].value))
    result = score_task(task, pred, True, tmp / "re")
    check("recalc pass all 5 cells", result.get("pass") is True, str(result))

    # without fill-down only B1 would match
    pred2 = tmp / "pred_nofill.xlsx"
    shutil.copy(init, pred2)
    wb2 = openpyxl.load_workbook(pred2)
    wb2["Sheet1"]["B1"] = "=A1*2"
    wb2.save(pred2)
    result2 = score_task(task, pred2, True, tmp / "re2")
    check("no fill-down fails (shows the feature matters)", result2.get("pass") is False, str(result2))

    # one formula across a row
    for c, n in zip("ABCDE", (1, 2, 3, 4, 5)):
        ws[f"{c}1"] = n
        ws[f"{c}2"] = n * 2
    wb.save(gold)
    for c in "ABCDE":
        ws[f"{c}2"] = None
    wb.save(init)
    row_task = {**task, "id": "synth-fill-right", "answer_position": "A2:E2"}
    write_output(row_task, SpreadsheetAnswer(cells=[CellValue(cell="A2", value="=A1*2")]), pred)
    written = openpyxl.load_workbook(pred)["Sheet1"]
    check("E2 got a formula", isinstance(written["E2"].value, str) and written["E2"].value.startswith("="), str(written["E2"].value))
    result_r = score_task(row_task, pred, True, tmp / "re-r")
    check("fill-right recalc pass", result_r.get("pass") is True, str(result_r))
    shutil.rmtree(tmp, ignore_errors=True)


def test_real_tasks_mocked() -> None:
    print("\n[5] real tasks, mocked model replies (no Tinker)")
    tasks = {str(t["id"]): t for t in load_dataset()}
    tmp = Path(tempfile.mkdtemp())

    # 51-12: one cell, thinking-wrapped JSON — parse path on a real workbook
    t = tasks["51-12"]
    out = tmp / "51-12.xlsx"
    reply = '<think>count increases... {not json}</think>\n```json\n{"cells":[{"cell":"B6","value":4}]}\n```'
    write_output(t, parse_answer(reply), out)
    result = score_task(t, out, bool(soffice_path()), tmp / "re51")
    check("51-12 thinking-wrapped JSON grades PASS", result.get("pass") is True, str(result))

    # 18645: two cells, literals still work after formula prompt change
    t2 = tasks.get("18645")
    if t2:
        from sb import load_answer_values

        gold = load_answer_values(t2["golden_xlsx"], t2)
        cells = [{"cell": coord, "value": v} for (_s, coord), v in gold.items()]
        reply2 = json.dumps({"cells": cells})
        out2 = tmp / "18645.xlsx"
        write_output(t2, parse_answer("Sure.\n" + reply2 + "\nhope that helps {ok}"), out2)
        result2 = score_task(t2, out2, bool(soffice_path()), tmp / "re186")
        check("18645 extra-brace JSON still writes gold", result2.get("pass") is True, str(result2))

    shutil.rmtree(tmp, ignore_errors=True)


def test_agent_loop() -> None:
    print("\n[6] python tool loop (no Tinker)")
    from agent import parse_agent_turn, run_python, should_use_agent

    kind, payload = parse_agent_turn('{"tool":"python","code":"print(1)"}')
    check("parse python tool", kind == "python" and "print" in payload)
    kind, _ = parse_agent_turn('{"tool":"done"}')
    check("parse done", kind == "done")
    kind, payload = parse_agent_turn('```python\nprint(wb.sheetnames)\n```')
    check("parse fence as python", kind == "python" and "sheetnames" in payload)

    tmp = Path(tempfile.mkdtemp())
    init = tmp / "init.xlsx"
    gold = tmp / "golden.xlsx"
    pred = tmp / "pred.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i, n in enumerate((1, 2, 3, 4, 5), start=1):
        ws[f"A{i}"] = n
        ws[f"B{i}"] = n * 2
    wb.save(gold)
    for i in range(1, 6):
        ws[f"B{i}"] = None
    wb.save(init)
    task = {
        "id": "synth-agent",
        "instruction": "double A into B",
        "init_xlsx": str(init),
        "golden_xlsx": str(gold),
        "answer_position": "B1:B5",
        "answer_sheet": "Sheet1",
        "instruction_type": "Cell-Level Manipulation",
    }
    live = openpyxl.load_workbook(init)
    result = run_python(
        "for sheet, coord in ANSWER:\n"
        "    ws = wb[sheet]\n"
        "    r = int(''.join(c for c in coord if c.isdigit()))\n"
        "    ws[coord] = ws[f'A{r}'].value * 2\n"
        "print('wrote', len(ANSWER))",
        live,
        task,
    )
    check("exec printed", "wrote 5" in result, result)
    live.save(pred)
    scored = score_task(task, pred, bool(soffice_path()), tmp / "re-ag")
    check("python loop fills B1:B5", scored.get("pass") is True, str(scored))
    safe = run_python(
        "import math\nprint(repr(type(1)), isinstance(1, int), hasattr(math, 'sqrt'))",
        live,
        task,
    )
    check("common builtins and safe import work", "<class 'int'>" in safe and "True True" in safe, safe)
    check("blocked import", "ERROR" in run_python("import os\nprint(os.getcwd())", live, task))
    dumped = run_python("print(sorted(task.keys())); print(task)", live, task)
    check("task has no golden path", "golden" not in dumped and "xlsx" not in dumped, dumped)
    check("task has no init path", "init_xlsx" not in dumped, dumped)
    small = {**task, "answer_position": "B1"}
    # should_use_agent loads init and counts B1:B5 from the real task; use a 1-cell task
    one = {**task, "answer_position": "B1"}
    check("small task stays one-shot", should_use_agent(one) is False)
    check("5-cell column uses one-shot", should_use_agent(task) is False)
    wide = {**task, "answer_position": "B1:B40"}
    check("40-cell column uses agent", should_use_agent(wide) is True)
    shutil.rmtree(tmp, ignore_errors=True)


def test_focus_and_thin() -> None:
    print("\n[7] focus windows + thin guard")
    tmp = Path(tempfile.mkdtemp())
    init = tmp / "init.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "top"
    ws["AE5"] = "far-col"
    ws["A150"] = "far-row"
    wb.save(init)
    task = {
        "id": "focus",
        "instruction": "x",
        "init_xlsx": str(init),
        "answer_position": "AE5",
        "answer_sheet": "Sheet1",
        "data_position": "A150",
    }
    dump = serialize_workbook(init, task=task)
    check("sheet name excludes display label", "### Sheet: Sheet1\nWindow: overview" in dump)
    check("overview still has A1", "top" in dump)
    check("focus includes far column", "far-col" in dump)
    check("focus includes far data row", "far-row" in dump)
    no_task = serialize_workbook(init)
    check("no-task crop hides AE5", "far-col" not in no_task)

    gold = tmp / "golden.xlsx"
    pred = tmp / "pred.xlsx"
    wb2 = openpyxl.Workbook()
    w2 = wb2.active
    w2.title = "Sheet1"
    for i in range(1, 25):
        w2[f"A{i}"] = i
    wb2.save(gold)
    for i in range(1, 25):
        w2[f"A{i}"] = None
    wb2.save(init)
    large = {
        "id": "thin-large",
        "instruction": "x",
        "init_xlsx": str(init),
        "golden_xlsx": str(gold),
        "answer_position": "A1:A24",
        "answer_sheet": "Sheet1",
    }
    from common import CellValue, SpreadsheetAnswer

    write_output(large, SpreadsheetAnswer(cells=[CellValue(cell="A1", value=1)]), pred)
    check("24-cell one-value is thin", output_is_thin(large, pred) is True)
    write_output(large, SpreadsheetAnswer(cells=[CellValue(cell="A1", value="=1")]), pred)
    # fill-down writes formulas on all 24
    check("24-cell one formula is not thin", output_is_thin(large, pred) is False)
    small = {**large, "answer_position": "A1:A7"}
    write_output(small, SpreadsheetAnswer(cells=[CellValue(cell="A1", value=1)]), pred)
    check("small task never thin", output_is_thin(small, pred) is False)

    two = tmp / "two.xlsx"
    wb3 = openpyxl.Workbook()
    a = wb3.active
    a.title = "PL"
    a["A1"] = None
    b = wb3.create_sheet("ST")
    b["A1"] = None
    wb3.save(two)
    multi = {
        "id": "two-sheet",
        "instruction": "x",
        "init_xlsx": str(two),
        "golden_xlsx": str(two),
        "answer_position": "PL'!A1,ST'!A1",
        "answer_sheet": "PL",
    }
    write_output(
        multi,
        SpreadsheetAnswer(cells=[CellValue(cell="A1", value="left"), CellValue(cell="A1", value="right")]),
        pred,
    )
    got = openpyxl.load_workbook(pred)
    check("multi-sheet A1 not collapsed", got["PL"]["A1"].value == "left" and got["ST"]["A1"].value == "right",
          f"PL={got['PL']['A1'].value!r} ST={got['ST']['A1'].value!r}")
    write_output(
        multi,
        SpreadsheetAnswer(cells=[
            CellValue(cell="ST!A1", value="right"),
            CellValue(cell="PL!A1", value="left"),
        ]),
        pred,
    )
    got2 = openpyxl.load_workbook(pred)
    check(
        "Sheet!A1 keys win even out of order",
        got2["PL"]["A1"].value == "left" and got2["ST"]["A1"].value == "right",
        f"PL={got2['PL']['A1'].value!r} ST={got2['ST']['A1'].value!r}",
    )
    got.close()
    got2.close()
    shutil.rmtree(tmp, ignore_errors=True)


def main() -> None:
    test_json_parse()
    test_xlfn()
    test_filldown()
    test_filldown_recalc()
    test_real_tasks_mocked()
    test_agent_loop()
    test_focus_and_thin()
    print(f"\n{PASS} passed, {FAIL} failed")
    raise SystemExit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
